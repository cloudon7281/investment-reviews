from datetime import datetime
import pandas as pd
import math
from typing import Dict, List, Optional, TypedDict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import locale
from numbers_parser import Document, Table, Style
from pathlib import Path
from enum import Enum
from logger import logger
from tabulate import tabulate
import math
import reporter_definitions as rd
import warnings
from data_table_builder import DataTableBuilder
from console_table_writer import ConsoleTableWriter
from numbers_table_writer import NumbersTableWriter
from csv_writer import CSVWriter

class CellColor(Enum):
    NONE = "none"
    RED = "red"
    AMBER = "amber"
    GREEN = "green"

class CellFormat(TypedDict, total=False):
    currency: str  # 3-letter currency code
    percentage: bool  # True for percentage formatting
    decimal_places: int  # Number of decimal places
    color: CellColor  # Cell background color

class PortfolioReporter:
    """Class for generating portfolio reports in Numbers format."""
    
    def __init__(self, numbers_filename: Optional[str] = None):
        """Initialize the reporter.
        
        Args:
            numbers_filename: Name of the Numbers file to create
        """
        self.numbers_filename = numbers_filename
        self.numbers_file = Path(numbers_filename) if numbers_filename else None
        self.numbers_doc = None
        self.numbers_row = 0
        
        # Initialize new infrastructure
        self.data_builder = DataTableBuilder()
        self.console_writer = ConsoleTableWriter()
        self.numbers_writer = None  # Will be initialized when Numbers doc is created
        
        # Set locale for currency formatting
        try:
            locale.setlocale(locale.LC_ALL, 'en_GB.UTF-8')
        except locale.Error:
            logger.warning("Could not set locale for currency formatting")
        
        if numbers_filename:
            # Set up the Numbers document and table
            self.numbers_doc = Document()
            self.numbers_writer = NumbersTableWriter(self.numbers_doc, numbers_filename)
        
    def _save_numbers_document(self) -> None:
        """Save the Numbers document to file."""
        if self.numbers_filename and self.numbers_doc:
            try:
                # Note: We can't easily remove sheets/tables with numbers_parser
                # The blank "Sheet 1" and "Table 1" will remain, but functionality works
                
                numbers_filename = str(Path(self.numbers_filename).with_suffix('.numbers'))
                self.numbers_doc.save(numbers_filename)
                logger.info(f"Saved Numbers report to {numbers_filename}")
            except Exception as e:
                logger.error(f"Error saving Numbers file: {str(e)}")
                logger.info("Continuing without Numbers output...")

    def display_full_history(self, full_history_results: Dict[str, pd.DataFrame]) -> None:
        """Display full history results using new clean infrastructure.
        
        Args:
            full_history_results: Dictionary with 'individual_stocks', 'whole_portfolio', and 'per_tag' DataFrames
        """
        logger.info("Displaying full history results")
        
        if not full_history_results or full_history_results.get('individual_stocks', pd.DataFrame()).empty:
            logger.warning("No full history data to display")
            return
        
        # 1. Combine whole portfolio, per-category, and per-tag data for summary table
        portfolio_summary_df = self._combine_portfolio_data(
            full_history_results['whole_portfolio'], 
            full_history_results.get('per_category', pd.DataFrame()),
            full_history_results['per_tag']
        )
        
        # 2. Prepare main stocks (just sorting)
        main_stocks_df = self._prepare_main_stocks(full_history_results['individual_stocks'])
        
        # 3. Define table sequence
        tables = [
            (portfolio_summary_df, 'tag_summary', 'Portfolio Summary'),
            (main_stocks_df, 'full_history', 'Full Investment History')
        ]
        
        # 4. Render all tables
        for df_table, config_name, title in tables:
            config = rd.COLUMN_CONFIGS[config_name]
            
            # Build table data using DataTableBuilder with title
            table_data = self.data_builder.build_table(df_table, config, title)
            
            # Console output
            self.console_writer.write_table(table_data, config)
            
            # Numbers output
            if self.numbers_writer:
                self.numbers_writer.write_table(table_data, config, title, title)
        
        # Save Numbers file if we have one
        if self.numbers_writer:
            self._save_numbers_document()
    
    def _combine_portfolio_data(self, whole_portfolio_df: pd.DataFrame, per_category_df: pd.DataFrame, per_tag_df: pd.DataFrame) -> pd.DataFrame:
        """Combine whole portfolio, per-category, and per-tag data for display.
        
        Args:
            whole_portfolio_df: DataFrame with single row for overall portfolio
            per_category_df: DataFrame with one row per account category (ISA/taxable/pension)
            per_tag_df: DataFrame with one row per tag
            
        Returns:
            Combined DataFrame ready for display in order: Whole Portfolio -> Per-Category -> Per-Tag
        """
        # Add row type markers before combining (for visual distinction)
        if not whole_portfolio_df.empty:
            whole_portfolio_df = whole_portfolio_df.copy()
            whole_portfolio_df['_row_type'] = 'portfolio'
        
        if not per_category_df.empty:
            per_category_df = per_category_df.copy()
            per_category_df['_row_type'] = 'category'
        
        if not per_tag_df.empty:
            per_tag_df = per_tag_df.copy()
            per_tag_df['_row_type'] = 'tag'
        
        # Combine the DataFrames in the correct order
        combined_df = pd.concat([whole_portfolio_df, per_category_df, per_tag_df], ignore_index=True)
        
        # Reorder to ensure correct hierarchy: Whole Portfolio -> Per-Category (by P&L) -> Per-Tag (by P&L)
        whole_portfolio_row = combined_df[combined_df['tag'] == 'Whole Portfolio']
        # Categories are now capitalized (ISA, Taxable, Pension)
        per_category_rows = combined_df[combined_df['tag'].isin(['ISA', 'Taxable', 'Pension'])]
        # Sort categories by P&L descending
        per_category_rows = per_category_rows.sort_values('total_pnl', ascending=False)
        
        # Per-tag rows are everything except Whole Portfolio and categories
        per_tag_rows = combined_df[~combined_df['tag'].isin(['Whole Portfolio', 'ISA', 'Taxable', 'Pension'])]
        # Sort tags by P&L descending
        per_tag_rows = per_tag_rows.sort_values('total_pnl', ascending=False)
        
        # Combine in the correct order
        combined_df = pd.concat([whole_portfolio_row, per_category_rows, per_tag_rows], ignore_index=True)
        
        return combined_df
    
    def _capitalize_category_display(self, df: pd.DataFrame) -> pd.DataFrame:
        """Capitalize category names for display.
        
        Args:
            df: DataFrame with 'tag' column containing category names
            
        Returns:
            DataFrame with capitalized category names for display
        """
        display_mapping = {
            'isa': 'ISA',
            'taxable': 'Taxable', 
            'pension': 'Pension'
        }
        df = df.copy()
        df['tag'] = df['tag'].map(display_mapping).fillna(df['tag'])
        return df
    
    def _prepare_main_stocks(self, df: pd.DataFrame) -> pd.DataFrame:
        """Prepare main stocks DataFrame with sorting and multi-category labeling.
        
        Args:
            df: DataFrame containing individual stock data (clean, no aggregations)
            
        Returns:
            DataFrame with stocks sorted by tag then P&L, with category suffix for multi-category stocks
        """
        # Detect tickers that appear in multiple categories
        multi_category_tickers = self._detect_multi_category_tickers(df)
        
        # Add category suffix to stock name for multi-category stocks
        sorted_df = df.copy()
        if multi_category_tickers:
            for idx, row in sorted_df.iterrows():
                if row['ticker'] in multi_category_tickers:
                    category_display = row['account_type'].upper() if row['account_type'] == 'isa' else row['account_type'].capitalize()
                    sorted_df.at[idx, 'stock_name'] = f"{row['stock_name']} ({category_display})"
        
        # Sort by tag first, then by P&L descending
        sorted_df['sort_tag'] = sorted_df['tag'].fillna('No Tag')
        sorted_df['sort_pnl'] = -sorted_df['total_pnl']  # Negative for descending order
        sorted_df = sorted_df.sort_values(['sort_tag', 'sort_pnl'])
        
        # Remove the temporary sort columns
        sorted_df = sorted_df.drop(['sort_tag', 'sort_pnl'], axis=1)
        
        return sorted_df
    
    def _detect_multi_category_tickers(self, df: pd.DataFrame) -> set:
        """Detect tickers that appear in multiple account categories.
        
        Args:
            df: DataFrame with 'ticker' and 'account_type' columns
            
        Returns:
            Set of tickers that appear in more than one category
        """
        ticker_categories = df.groupby('ticker')['account_type'].nunique()
        multi_category_tickers = set(ticker_categories[ticker_categories > 1].index)
        return multi_category_tickers

    def display_periodic_review(self, periodic_results: Dict[str, pd.DataFrame], start_date: datetime, end_date: datetime, eval_date: Optional[datetime] = None) -> None:
        """Display periodic review using new clean infrastructure.
        
        Args:
            periodic_results: Dictionary with periodic review data
            start_date: Start of the review period
            end_date: End of the review period
            eval_date: Evaluation date (defaults to now if None)
        """
        logger.info("Displaying periodic review results")
        
        if eval_date is None:
            eval_date = datetime.now()
        
        # 1. Combine summary and per_tag data for summary table (like full-history)
        summary_df = self._combine_periodic_summary_data(
            periodic_results.get('summary', pd.DataFrame()),
            periodic_results.get('per_tag', pd.DataFrame())
        )
        
        # 2. Prepare detail DataFrames (just sorting)
        new_df = self._prepare_periodic_detail(periodic_results.get('new', pd.DataFrame()))
        retained_df = self._prepare_periodic_detail(periodic_results.get('retained', pd.DataFrame()))
        sold_df = self._prepare_periodic_detail(periodic_results.get('sold', pd.DataFrame()))
        
        # 3. Define table sequence
        tables = []
        
        # Summary table (combined summary + per_tag)
        if not summary_df.empty:
            tables.append((summary_df, 'periodic_review_summary', f'Periodic Review Summary ({start_date.strftime("%Y-%m-%d")} to {end_date.strftime("%Y-%m-%d")}, evaluated on {eval_date.strftime("%Y-%m-%d")})'))
        
        # Detail tables
        for category, df in [('New', new_df), ('Retained', retained_df), ('Sold', sold_df)]:
            if not df.empty:
                tables.append((df, 'periodic_review_detail', f'{category} Stocks ({start_date.strftime("%Y-%m-%d")} to {end_date.strftime("%Y-%m-%d")}, evaluated on {eval_date.strftime("%Y-%m-%d")})'))
        
        # 4. Render all tables
        for df_table, config_name, title in tables:
            config = rd.COLUMN_CONFIGS[config_name]
            
            # Build table data using DataTableBuilder with title
            table_data = self.data_builder.build_table(df_table, config, title)
            
            # Console output
            self.console_writer.write_table(table_data, config)
            
            # Numbers output
            if self.numbers_writer:
                self.numbers_writer.write_table(table_data, config, title, title)
        
        # Save Numbers file if we have one
        if self.numbers_writer:
            self._save_numbers_document()
    
    def _combine_periodic_summary_data(self, summary_df: pd.DataFrame, per_tag_df: pd.DataFrame) -> pd.DataFrame:
        """Combine periodic review summary and per_tag data for display.
        
        Args:
            summary_df: DataFrame with overall summary (New/Retained/Sold totals)
            per_tag_df: DataFrame with per-tag summary data
            
        Returns:
            Combined DataFrame ready for display
        """
        # Separate summary rows (keep at top)
        summary_rows = summary_df.copy()
        
        # Sort per_tag data by category order (new, retained, sold) then by P&L descending within each category
        if not per_tag_df.empty:
            # Extract P&L value for sorting
            per_tag_df['_sort_pnl'] = per_tag_df['pnl'].apply(lambda x: x[0] if isinstance(x, tuple) else x)
            
            # Define category order
            category_order = {'new': 0, 'retained': 1, 'sold': 2}
            per_tag_df['_category_order'] = per_tag_df['sort_category'].map(category_order)
            
            # Sort by category order first, then by P&L descending within each category
            per_tag_df = per_tag_df.sort_values(['_category_order', '_sort_pnl'], ascending=[True, False])
            
            # Drop the temporary sorting columns
            per_tag_df = per_tag_df.drop(['_sort_pnl', '_category_order', 'sort_category', 'sort_pnl'], axis=1, errors='ignore')
        
        # Combine summary rows at top, then tag rows
        combined_df = pd.concat([summary_rows, per_tag_df], ignore_index=True)
        
        return combined_df
    
    def _prepare_periodic_detail(self, df: pd.DataFrame) -> pd.DataFrame:
        """Prepare periodic review detail DataFrame.
        
        Args:
            df: DataFrame containing detail data
            
        Returns:
            DataFrame ready for display
        """
        if df.empty:
            return df

        # Sort by tag first, then by P&L descending (same as console output)
        sorted_df = df.copy()
        sorted_df['sort_tag'] = sorted_df['tag'].fillna('No Tag')
        
        # Sort by P&L (extract value from tuple)
        sorted_df['sort_pnl'] = -sorted_df['pnl'].apply(lambda x: x[0] if isinstance(x, tuple) else x)  # Negative for descending order
        
        sorted_df = sorted_df.sort_values(['sort_tag', 'sort_pnl'])
        
        # Remove the temporary sort columns
        sorted_df = sorted_df.drop(['sort_tag', 'sort_pnl'], axis=1)
        
        return sorted_df
    
    def display_tax_report(self, tax_report_results: Dict, tax_year: str):
        """Display tax report using new clean infrastructure.
        
        Args:
            tax_report_results: Dict with 'summary' and 'transactions' DataFrames
            tax_year: Tax year string (e.g., 'FY25')
        """
        logger.info(f"Displaying tax report for {tax_year}")
        
        summary_df = tax_report_results.get('summary', pd.DataFrame())
        transactions_df = tax_report_results.get('transactions', pd.DataFrame())
        
        if transactions_df.empty:
            print(f"\n{'='*80}")
            print(f"TAX REPORT FOR {tax_year.upper()}")
            print(f"{'='*80}")
            print("No taxable transactions found for this tax year.")
            return
        
        # Prepare tables
        summary_table = self._prepare_tax_report_summary(summary_df, tax_year)
        transactions_table = self._prepare_tax_report_transactions(transactions_df)
        
        # Create table builders and writers
        data_builder = DataTableBuilder()
        console_writer = ConsoleTableWriter()
        numbers_writer = self.numbers_writer if self.numbers_doc else None
        
        # Process tables
        tables = [
            (summary_table, 'tax_report_summary', "Tax Report Summary"),
            (transactions_table, 'tax_report', "Taxable Transactions")
        ]
        
        for df, config_name, title in tables:
            if df.empty:
                continue
                
            # Get configuration
            config = rd.COLUMN_CONFIGS[config_name]
            
            # Build table data
            table_data = data_builder.build_table(df, config, title=title)
            
            # Render to console
            console_writer.write_table(table_data, config)
            
            # Render to Numbers
            if numbers_writer:
                numbers_writer.write_table(table_data, config, title, title)
        
        # Save Numbers file
        if self.numbers_filename:
            self._save_numbers_document()
    
    def _prepare_tax_report_summary(self, summary_df: pd.DataFrame, tax_year: str) -> pd.DataFrame:
        """Prepare tax report summary data for display."""
        if summary_df.empty:
            return pd.DataFrame()
        
        row = summary_df.iloc[0]
        
        # Create a single row with properly typed columns
        summary_data = {
            'tax_year': tax_year,
            'total_transactions': row['total_transactions'],
            'net_gains_losses': row['net_gains_losses']
        }
        
        return pd.DataFrame([summary_data])
    
    def _prepare_tax_report_transactions(self, transactions_df: pd.DataFrame) -> pd.DataFrame:
        """Prepare tax report transactions data for display."""
        if transactions_df.empty:
            return pd.DataFrame()
        
        # Convert monetary columns to tuples if they aren't already
        display_df = transactions_df.copy()
        
        # Convert monetary columns to tuples if they aren't already
        for col in ['amount_received', 'total_price_paid', 'average_price', 'pnl']:
            if col in display_df.columns:
                # Check if any value in the column is already a tuple
                has_tuples = any(isinstance(val, tuple) for val in display_df[col] if pd.notna(val))
                if not has_tuples:
                    display_df[col] = display_df[col].apply(lambda x: (x, 'GBP') if pd.notna(x) else (0.0, 'GBP'))
        
        return display_df
    
    def write_value_over_time_csv(self, value_df: pd.DataFrame, n_days: int) -> None:
        """Write value-over-time CSV file.

        Args:
            value_df: DataFrame with value-over-time data from PortfolioAnalysis
            n_days: Number of days (for logging purposes)
        """
        if not self.numbers_filename:
            logger.warning("Cannot write value-over-time CSV: no output file specified")
            return

        if value_df.empty:
            logger.warning("No value-over-time data to write")
            return

        logger.info(f"Writing value-over-time CSV for {n_days} days")

        # Derive CSV filename from Numbers filename
        csv_filename = str(Path(self.numbers_filename).with_suffix('')) + '_value_over_time.csv'

        # Create CSV writer and write the file
        csv_writer = CSVWriter(csv_filename)
        csv_writer.write_value_over_time(value_df)

        logger.info(f"Value-over-time CSV written: {csv_filename}")

    def display_annual_review(self, annual_results: Dict[str, pd.DataFrame], start_date: datetime) -> None:
        """Display annual review results.

        Args:
            annual_results: Dictionary with annual review data:
                - 'whole_portfolio': DataFrame with single summary row
                - 'per_category': DataFrame with ISA/Taxable/Pension rows
                - 'per_tag': DataFrame with per-tag summaries
                - 'individual_stocks': DataFrame with per-stock detail
            start_date: Start date of the review period
        """
        logger.info("Displaying annual review results")

        eval_date = datetime.now()

        # 1. Combine summary data (whole portfolio, per-category, per-tag)
        summary_df = self._combine_annual_summary_data(
            annual_results.get('whole_portfolio', pd.DataFrame()),
            annual_results.get('per_category', pd.DataFrame()),
            annual_results.get('per_tag', pd.DataFrame())
        )

        # 2. Prepare individual stocks (sorting)
        individual_df = self._prepare_annual_detail(annual_results.get('individual_stocks', pd.DataFrame()))

        # 3. Define table sequence
        tables = []

        # Summary table
        if not summary_df.empty:
            title = f'Annual Review Summary ({start_date.strftime("%Y-%m-%d")} to {eval_date.strftime("%Y-%m-%d")})'
            tables.append((summary_df, 'annual_review_summary', title))

        # Detail table
        if not individual_df.empty:
            title = f'Annual Review Detail ({start_date.strftime("%Y-%m-%d")} to {eval_date.strftime("%Y-%m-%d")})'
            tables.append((individual_df, 'annual_review_detail', title))

        # 4. Render all tables
        for df_table, config_name, title in tables:
            config = rd.COLUMN_CONFIGS[config_name]

            # Build table data using DataTableBuilder with title
            table_data = self.data_builder.build_table(df_table, config, title)

            # Console output
            self.console_writer.write_table(table_data, config)

            # Numbers output
            if self.numbers_writer:
                self.numbers_writer.write_table(table_data, config, title, title)

        # Save Numbers file if we have one
        if self.numbers_writer:
            self._save_numbers_document()

    def _combine_annual_summary_data(self, whole_portfolio_df: pd.DataFrame,
                                      per_category_df: pd.DataFrame,
                                      per_tag_df: pd.DataFrame) -> pd.DataFrame:
        """Combine annual review summary data for display.

        Args:
            whole_portfolio_df: DataFrame with single row for overall portfolio
            per_category_df: DataFrame with one row per account category
            per_tag_df: DataFrame with one row per tag

        Returns:
            Combined DataFrame in order: Whole Portfolio -> Per-Category -> Per-Tag
        """
        # Combine the DataFrames in the correct order
        combined_df = pd.concat([whole_portfolio_df, per_category_df, per_tag_df], ignore_index=True)

        if combined_df.empty:
            return combined_df

        # Reorder to ensure correct hierarchy
        whole_portfolio_row = combined_df[combined_df['group'] == 'Whole Portfolio']
        per_category_rows = combined_df[combined_df['group'].isin(['ISA', 'Taxable', 'Pension'])]
        per_tag_rows = combined_df[~combined_df['group'].isin(['Whole Portfolio', 'ISA', 'Taxable', 'Pension'])]

        # Sort categories and tags by P&L descending
        if not per_category_rows.empty:
            per_category_rows = per_category_rows.sort_values('pnl', ascending=False)
        if not per_tag_rows.empty:
            per_tag_rows = per_tag_rows.sort_values('pnl', ascending=False)

        # Combine in the correct order
        combined_df = pd.concat([whole_portfolio_row, per_category_rows, per_tag_rows], ignore_index=True)

        return combined_df

    def _prepare_annual_detail(self, df: pd.DataFrame) -> pd.DataFrame:
        """Prepare annual review detail DataFrame.

        Args:
            df: DataFrame containing individual stock data

        Returns:
            DataFrame sorted by tag then P&L
        """
        if df.empty:
            return df

        # Detect tickers that appear in multiple categories
        multi_category_tickers = self._detect_multi_category_tickers(df)

        # Add category suffix to stock name for multi-category stocks
        sorted_df = df.copy()
        if multi_category_tickers:
            for idx, row in sorted_df.iterrows():
                if row['ticker'] in multi_category_tickers:
                    category_display = row['account_type']
                    sorted_df.at[idx, 'stock_name'] = f"{row['stock_name']} ({category_display})"

        # Sort by tag first, then by P&L descending, then by ticker for stability
        sorted_df['sort_tag'] = sorted_df['tag'].fillna('No Tag')
        sorted_df['sort_pnl'] = -sorted_df['pnl']  # Negative for descending order
        sorted_df = sorted_df.sort_values(['sort_tag', 'sort_pnl', 'ticker'])

        # Remove the temporary sort columns
        sorted_df = sorted_df.drop(['sort_tag', 'sort_pnl'], axis=1)

        return sorted_df

    def write_price_over_time_csv(self, price_df: pd.DataFrame, start_date: datetime) -> None:
        """Write price-over-time CSV file.

        Args:
            price_df: DataFrame with price-over-time data (columns: date, ticker1, ticker2, ...)
            start_date: Start date of the period (for logging)
        """
        if not self.numbers_filename:
            logger.warning("Cannot write price-over-time CSV: no output file specified")
            return

        if price_df is None or price_df.empty:
            logger.warning("No price-over-time data to write")
            return

        logger.info(f"Writing price-over-time CSV from {start_date.strftime('%Y-%m-%d')}")

        # Derive CSV filename from Numbers filename
        csv_filename = str(Path(self.numbers_filename).with_suffix('')) + '_price_over_time.csv'

        # Create CSV writer and write the file
        csv_writer = CSVWriter(csv_filename)
        csv_writer.write_price_over_time(price_df)

        logger.info(f"Price-over-time CSV written: {csv_filename}")